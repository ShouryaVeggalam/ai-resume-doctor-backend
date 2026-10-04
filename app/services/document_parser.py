from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader


ALLOWED_EXTENSIONS = {".pdf", ".docx"}


async def extract_text(file: UploadFile, max_size_mb: int) -> str:
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Only PDF and DOCX resumes are supported.")

    raw = await file.read()
    max_bytes = max_size_mb * 1024 * 1024
    if len(raw) > max_bytes:
        raise HTTPException(status_code=413, detail=f"Resume exceeds the {max_size_mb} MB limit.")
    if not raw:
        raise HTTPException(status_code=400, detail="The uploaded resume is empty.")

    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(raw))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        else:
            doc = Document(BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not parse the resume file.") from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) < 40:
        raise HTTPException(status_code=400, detail="Not enough readable text was found in the resume.")
    return text[:120_000]
